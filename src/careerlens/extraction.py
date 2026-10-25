"""Local extraction for plain text, PDF, and DOCX documents."""

from pathlib import Path
import re

MAX_BYTES = 10 * 1024 * 1024


def normalize(text: str) -> str:
    return re.sub(
        r"[ \t]+", " ", re.sub(r"\n{3,}", "\n\n", text.replace("\r\n", "\n").replace("\r", "\n"))
    ).strip()


def extract_upload(name: str, data: bytes) -> str:
    if len(data) > MAX_BYTES:
        raise ValueError("File exceeds the 10 MB upload limit.")
    suffix = Path(name).suffix.lower()
    if suffix == ".txt":
        text = data.decode("utf-8-sig")
    elif suffix == ".pdf":
        import fitz

        with fitz.open(stream=data, filetype="pdf") as doc:
            text = "\n".join(page.get_text() for page in doc)
        if not text.strip():
            raise ValueError("No selectable text found. This may be an image-only scan; OCR is not included.")
    elif suffix == ".docx":
        from docx import Document
        import io

        doc = Document(io.BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs)
        text += "\n" + "\n".join(
            " | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows
        )
    else:
        raise ValueError("Supported formats are PDF, DOCX, and TXT.")
    if not text.strip():
        raise ValueError("The document contains no extractable text.")
    return normalize(text)

